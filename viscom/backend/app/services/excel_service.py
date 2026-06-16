import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _thin_border():
    side = Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)


def generate_report_excel(title: str, headers: list[str], rows: list[list], totals: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=14)

    ws["A1"] = title
    ws["A1"].font = title_font
    ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = _thin_border()

    for row_idx, row in enumerate(rows, 3):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = _thin_border()
            if isinstance(value, float | int) and not isinstance(value, bool):
                cell.alignment = Alignment(horizontal="right")

    if totals:
        total_row = len(rows) + 3
        for col_idx, header in enumerate(headers, 1):
            key = header.lower().replace(" ", "_")
            if key in totals:
                cell = ws.cell(row=total_row, column=col_idx, value=totals[key])
                cell.font = Font(bold=True)
                cell.border = _thin_border()
            elif col_idx == 1:
                cell = ws.cell(row=total_row, column=col_idx, value="TOTAL")
                cell.font = Font(bold=True)
                cell.border = _thin_border()

    for col in ws.columns:
        max_length = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_length + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
